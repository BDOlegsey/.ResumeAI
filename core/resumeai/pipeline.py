import io
import re
import logging
import zipfile
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from .models import AgentState, UserData
from .agents import search_agent_node, generator_agent_node, checker_agent_node
from .database import init_db

logger = logging.getLogger(__name__)


def parse_and_add_text(paragraph, text: str):
    """
    Парсит текст с Markdown-разметкой и добавляет в параграф с правильным форматированием
    Поддерживает: **жирный**, *курсив*, `код`
    """
    # Паттерны для разметки
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue

        # Жирный текст: **текст**
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        # Курсив: *текст*
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        # Код: `текст`
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        # Обычный текст
        else:
            paragraph.add_run(part)


def add_formatted_paragraph(doc: Document, text: str, style=None):
    """Добавляет параграф с поддержкой Markdown-форматирования"""
    text = text.rstrip()

    # Заголовки
    if text.startswith('### '):
        doc.add_heading(text[4:], level=3)
        return
    if text.startswith('## '):
        doc.add_heading(text[3:], level=2)
        return
    if text.startswith('# '):
        doc.add_heading(text[2:], level=1)
        return

    # Списки
    if text.strip().startswith('- ') or text.strip().startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        parse_and_add_text(p, text.strip()[2:])
        return

    # Нумерованные списки
    match = re.match(r'^(\d+)\.\s+(.+)$', text.strip())
    if match:
        p = doc.add_paragraph(style='List Number')
        parse_and_add_text(p, match.group(2))
        return

    # Обычный параграф
    p = doc.add_paragraph(style=style)
    parse_and_add_text(p, text)


def create_single_resume_docx(job_title: str, employer: str, resume_text: str) -> bytes:
    """Создаёт один DOCX-файл для одной компании"""
    doc = Document()

    # Заголовок
    title = doc.add_heading(f'Резюме: {job_title}', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Подзаголовок с компанией
    subtitle = doc.add_paragraph(f'Для компании: {employer}')
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(0, 102, 204)
    subtitle_run.bold = True

    doc.add_paragraph()  # Пустая строка

    # Основной текст резюме
    lines = resume_text.split('\n')
    for line in lines:
        if line.strip():
            add_formatted_paragraph(doc, line)

    # Сохранение в BytesIO
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_resume_for_employer(user_data: UserData, employer: str, max_iterations: int = 3) -> str:
    """Генерирует резюме для одной конкретной компании"""
    logger.info(f"Generating resume for: {employer}")

    employer_data = UserData(
        employers=[employer],
        job_title=user_data.job_title,
        experience=user_data.experience,
        skills=user_data.skills,
        achievements=user_data.achievements,
        attachments=user_data.attachments
    )

    state = AgentState(user_data=employer_data, max_iterations=max_iterations, current_employer=employer)

    # Поиск информации
    state = search_agent_node(state)
    if state.error_message:
        logger.error(f"Search failed for {employer}: {state.error_message}")
        return f"[Ошибка поиска информации для {employer}]"

    # Генерация и валидация
    while state.iteration_count < state.max_iterations:
        state = generator_agent_node(state)
        if state.error_message:
            logger.error(f"Generation failed for {employer}: {state.error_message}")
            return f"[Ошибка генерации резюме для {employer}]"

        state = checker_agent_node(state)
        if state.error_message:
            logger.warning(f"Checker failed for {employer}, using draft anyway")
            break

        if state.resume_draft and state.resume_draft.is_validated:
            logger.info(f"Resume validated for {employer}")
            break

    if not state.resume_draft:
        return f"[Не удалось сгенерировать резюме для {employer}]"

    return state.resume_draft.draft_content


def generate_resume_docx(user_data: UserData, max_iterations: int = 3) -> tuple[str, bytes, dict]:
    """
    Генерирует резюме для ВСЕХ компаний
    Возвращает: (объединённый_текст, zip_bytes, {employer: docx_bytes})
    """
    init_db()

    if not user_data.employers:
        raise ValueError("Список работодателей пуст")

    logger.info(f"Generating resumes for {len(user_data.employers)} employer(s)")

    employers_resumes = {}  # {employer: resume_text}
    employers_docx = {}  # {employer: docx_bytes}
    all_texts = []

    # Генерируем резюме для каждой компании
    for employer in user_data.employers:
        logger.info(f"Processing: {employer}")
        resume_text = generate_resume_for_employer(user_data, employer, max_iterations)
        employers_resumes[employer] = resume_text
        all_texts.append(f"=== РЕЗЮМЕ ДЛЯ: {employer.upper()} ===\n\n{resume_text}")

        # Создаём отдельный DOCX для этой компании
        docx_bytes = create_single_resume_docx(user_data.job_title, employer, resume_text)
        employers_docx[employer] = docx_bytes

    # Объединяем все тексты
    combined_text = "\n\n" + "=" * 80 + "\n\n".join(all_texts)

    # Создаём ZIP-архив со всеми резюме
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for employer, docx_bytes in employers_docx.items():
            # Безопасное имя файла (убираем спецсимволы)
            safe_name = re.sub(r'[^\w\s-]', '', employer).strip().replace(' ', '_')
            filename = f"resume_{safe_name}.docx"
            zip_file.writestr(filename, docx_bytes)

    zip_buffer.seek(0)
    zip_bytes = zip_buffer.getvalue()

    logger.info(f"Successfully generated {len(employers_docx)} resumes")
    return combined_text, zip_bytes, employers_docx
