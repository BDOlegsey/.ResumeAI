import logging
import os
import re
from docx import Document
from langgraph.graph import StateGraph, END
from models import AgentState, UserData
from agents import main_agent_node, search_agent_node, generator_agent_node, checker_agent_node
from database import init_db
from config import OUTPUT_DIR

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)

# --- Простейший разбор Markdown для жирного/заголовков/буллитов ---
BOLD_PAT = re.compile(r'(\*\*.*?\*\*)')

def add_markdown_line(doc: Document, line: str):
    """Добавляет строку с поддержкой #/##, - и **bold** в docx."""
    s = line.rstrip()

    # Заголовки
    if s.startswith("## "):
        doc.add_heading(s[3:], level=2)
        return
    if s.startswith("# "):
        doc.add_heading(s[2:], level=1)
        return

    # Буллиты
    if s.strip().startswith("- "):
        doc.add_paragraph(s.strip()[2:], style='List Bullet')
        return

    # Обычный параграф с **bold**
    p = doc.add_paragraph()
    parts = BOLD_PAT.split(s)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            text = part[2:-2]
            run = p.add_run(text)
            run.bold = True  # жирный на уровне Run
        else:
            p.add_run(part)

def save_resume_node(state: AgentState) -> AgentState:
    logger.info(f"Saving validated resume for {state.current_employer}.")
    if not state.resume_draft or not state.resume_draft.is_validated:
        logger.error("Save node called but resume draft is not validated.")
        state.error_message = "Пытались сохранить невалидированное резюме."
        return state

    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR, exist_ok=True)

    safe_employer = (state.current_employer or "UnknownEmployer").replace(' ', '_')
    safe_title = (state.user_data.job_title if state.user_data else "UnknownTitle").replace(' ', '_')
    filepath = os.path.join(OUTPUT_DIR, f"resume_{safe_employer}_{safe_title}.docx")

    try:
        doc = Document()
        # Заголовок файла — русскоязычный
        doc.add_heading(f"Резюме: {state.user_data.job_title} — {state.current_employer}", 0)
        for line in state.resume_draft.draft_content.split('\n'):
            if line.strip():
                add_markdown_line(doc, line)
        doc.save(filepath)
        logger.info(f"Resume saved to {filepath}")
        state.output_file_path = filepath
    except Exception as e:
        logger.error(f"Failed to save resume for {state.current_employer}: {e}")
        state.error_message = f"Не удалось сохранить резюме: {e}"
    return state

def main():
    init_db()

    workflow = StateGraph(AgentState)
    workflow.add_node("main_agent", main_agent_node)
    workflow.add_node("search_agent", search_agent_node)
    workflow.add_node("generator_agent", generator_agent_node)
    workflow.add_node("checker_agent", checker_agent_node)
    workflow.add_node("saver_node", save_resume_node)

    workflow.add_edge("main_agent", "search_agent")
    workflow.add_edge("search_agent", "generator_agent")
    workflow.add_edge("generator_agent", "checker_agent")

    def should_continue_validation(state: AgentState):
        if state.error_message:
            logger.error(f"Workflow stopped due to error: {state.error_message}")
            return "error"
        if state.resume_draft and state.resume_draft.is_validated:
            return "validated"
        elif state.resume_draft and state.iteration_count >= state.max_iterations:
            logger.warning(f"Max iterations ({state.max_iterations}) reached for {state.current_employer}.")
            return "validated"
        else:
            return "not_validated"

    workflow.add_conditional_edges(
        "checker_agent",
        should_continue_validation,
        {"validated": "saver_node", "not_validated": "generator_agent", "error": END}
    )

    def should_process_next_employer(state: AgentState):
        if state.error_message:
            return "error"
        if state.user_data and state.current_employer == state.user_data.employers[-1]:
            return "end"
        else:
            return "continue"

    workflow.add_conditional_edges(
        "saver_node",
        should_process_next_employer,
        {"continue": "main_agent", "end": END, "error": END}
    )

    workflow.set_entry_point("main_agent")
    app = workflow.compile()

    # Пример русскоязычного ввода
    user_input_data = {
        "user_data": {
            "employers": ["Yandex", "АО Элара", "Букет Чувашии"],
            "job_title": "Senior Software Engineer",
            "experience": (
                "Продуктовая IT‑компания (NDA) — Senior Software Engineer (03.2022–настоящее время) [web:51]\n"
                "• Проектирование и развитие микросервисов (Python/FastAPI, Node.js), интеграции с внешними API; лидирование команды из 3 разработчиков [web:21]\n"
                "• Внедрение CI/CD (Docker, GitLab CI/GitHub Actions), автоматизация тестов; мониторинг и алертинг (Prometheus, Grafana) [web:21]\n"
                "• Результаты: время загрузки интерфейса −40%, частота релизов +50%, p95 задержка API −35%, MTTR −30% [web:28][web:33]\n"
                "\n"
                "Системный интегратор — Software Engineer (06.2020–02.2022) [web:47]\n"
                "• Разработка высоконагруженных REST/GraphQL API, оптимизация SQL‑запросов (PostgreSQL), кеширование (Redis), фоновые задачи (Celery) [web:21]\n"
                "• Практики качества: code review, pytest, статический анализ, контроль производительности и безопасности [web:21]\n"
                "\n"
                "Yamate — Junior Software Engineer (08.2018–05.2020) [web:47]\n"
                "• Внутренние веб‑панели (React) и сервисы, контейнеризация в Docker, стандартизация деплоя и автоматизация отчётности [web:21]\n"
            ),
            "skills": "Python (FastAPI, Django, Flask), JavaScript/TypeScript (React, Node.js), PostgreSQL, Redis, Docker, AWS, Git, Nginx, Prometheus, Grafana, pytest, Agile/Scrum [web:21]",
            "achievements": (
                "• Сократил время загрузки интерфейса на 40% (оптимизация бандла, кеширование, lazy loading) [web:28][web:33]\n"
                "• Внедрил CI/CD и автотесты, увеличив частоту релизов на 50% при сохранении стабильности [web:28][web:33]\n"
                "• Снизил p95 задержку API на 35% за счёт оптимизации запросов и кеширования [web:28][web:33]\n"
                "• Сократил MTTR на 30% благодаря наблюдаемости и корректному алертингу [web:28][web:33]\n"
            ),
            "attachments": []
        }
    }

    initial_state = AgentState(user_data=UserData(**user_input_data["user_data"]))

    logger.info("Starting the multi-agent resume generation workflow.")
    try:
        final_state = app.invoke(initial_state)
        logger.info("Workflow completed.")
        if isinstance(final_state, dict):
            err = final_state.get("error_message")
            if err:
                logger.error(f"Final state error: {err}")
            else:
                logger.info(f"Resumes generated successfully. Path: {final_state.get('output_file_path')}")
        else:
            if getattr(final_state, "error_message", None):
                logger.error(f"Final state error: {final_state.error_message}")
            else:
                logger.info(f"Resumes generated successfully. Path: {getattr(final_state, 'output_file_path', None)}")
    except Exception as e:
        logger.error(f"Workflow execution failed: {e}")

if __name__ == "__main__":
    main()
