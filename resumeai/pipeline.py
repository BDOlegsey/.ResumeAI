import io
import re
import logging
from docx import Document
from .models import AgentState, UserData
from .agents import main_agent_node, search_agent_node, generator_agent_node, checker_agent_node
from .database import init_db

logger = logging.getLogger(__name__)
BOLD_PAT = re.compile(r'(\*\*.*?\*\*)')

def _add_markdown_line(doc: Document, line: str):
    s = line.rstrip()
    if s.startswith("## "):
        doc.add_heading(s[3:], level=2)
        return
    if s.startswith("# "):
        doc.add_heading(s[2:], level=1)
        return
    if s.strip().startswith("- "):
        doc.add_paragraph(s.strip()[2:], style='List Bullet')
        return
    p = doc.add_paragraph()
    parts = BOLD_PAT.split(s)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)

def build_docx_from_text(title: str, text: str) -> bytes:
    doc = Document()
    if title:
        doc.add_heading(title, 0)
    for line in text.split("\n"):
        if line.strip():
            _add_markdown_line(doc, line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def generate_resume_docx(user_data: UserData, max_iterations: int = 3) -> tuple[str, bytes]:
    init_db()
    state = AgentState(user_data=user_data, max_iterations=max_iterations)

    state = main_agent_node(state)
    if state.error_message:
        raise RuntimeError(state.error_message)

    state = search_agent_node(state)
    if state.error_message:
        raise RuntimeError(state.error_message)

    while state.iteration_count < state.max_iterations:
        state = generator_agent_node(state)
        if state.error_message:
            raise RuntimeError(state.error_message)
        state = checker_agent_node(state)
        if state.error_message:
            raise RuntimeError(state.error_message)
        if state.resume_draft and state.resume_draft.is_validated:
            break

    if not state.resume_draft:
        raise RuntimeError("Не удалось сгенерировать резюме")

    title = f"Резюме: {user_data.job_title} — {state.current_employer or ''}".strip()
    docx_bytes = build_docx_from_text(title, state.resume_draft.draft_content)
    return state.resume_draft.draft_content, docx_bytes
